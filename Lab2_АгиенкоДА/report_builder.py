"""Сборка отчёта по лабораторной работе №2 (формат .docx)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# --- Титул -----------------------------------------------------------
title = doc.add_heading("Отчёт по лабораторной работе №2", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("Тема: Юнит-тестирование проектов и автоматизация проверок")
doc.add_paragraph("Студент: Агиенко Данил Алексеевич")
doc.add_paragraph("Группа: 2307В2")
doc.add_paragraph("Дисциплина: PTPM-VKI (программное тестирование)")

doc.add_paragraph()

# --- Пункт А ---------------------------------------------------------
doc.add_heading("Пункт А. Общее количество тестов и статистика прохождения", level=1)

table = doc.add_table(rows=3, cols=4)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Модуль"
hdr[1].text = "Всего тестов"
hdr[2].text = "Прошло (ok)"
hdr[3].text = "Упало (fail/error)"

row = table.rows[1].cells
row[0].text = "Свой проект (my_project.py)"
row[1].text = "28"
row[2].text = "28"
row[3].text = "0"

row = table.rows[2].cells
row[0].text = "Модуль доставки (delivery_service.py)"
row[1].text = "26"
row[2].text = "24"
row[3].text = "2"

doc.add_paragraph()
doc.add_paragraph(
    "Итого написано 54 юнит-теста (минимальное требование — 20). "
    "Тесты собственного проекта после исправления бага mask_password "
    "проходят полностью (28/28). По модулю доставки выявлено 2 дефекта."
)

# --- Пункт Б ---------------------------------------------------------
doc.add_heading("Пункт Б. Количество упавших тестов для кода доставки", level=1)

p = doc.add_paragraph()
run = p.add_run("Упало 2 теста из 26 (8 %).")
run.bold = True

doc.add_paragraph(
    "Тесты падают не из-за ошибок в самих тестах, а из-за дефектов "
    "в логике исходного кода delivery_service.py."
)

# --- Пункт В ---------------------------------------------------------
doc.add_heading("Пункт В. Локализация аномалий (дефектов)", level=1)

doc.add_paragraph(
    "Ниже описаны оба найденных дефекта: название теста, номер строки, "
    "описание некорректного поведения и предлагаемое исправление."
)

# Дефект 1
doc.add_heading("Дефект №1. Экспресс-доставка стоит дешевле обычной", level=2)

t1 = doc.add_table(rows=6, cols=2)
t1.style = "Light Grid Accent 1"
cells = t1.rows[0].cells
cells[0].text = "Параметр"
cells[1].text = "Значение"

rows_data_1 = [
    ("Упавший тест", "test_express_is_more_expensive_than_regular"),
    ("Номер строки", "36 (файл delivery_service.py)"),
    ("Код строки", "if is_express: total_cost *= 0.5"),
    ("Описание дефекта",
     "Логика экспресс-доставки умножает стоимость на 0.5, т.е. делает "
     "ускоренную доставку ВДВОЕ ДЕШЕВЛЕ обычной. По бизнес-смыслу экспресс — "
     "это дополнительная услуга, которая должна стоить ДОРОЖЕ обычной доставки."),
    ("Вариант исправления",
     "Заменить на надбавку, например: if is_express: total_cost *= 1.5 "
     "(или total_cost += фиксированная наценка)."),
]
for i, (k, v) in enumerate(rows_data_1, start=1):
    cells = t1.rows[i].cells
    cells[0].text = k
    cells[1].text = v

doc.add_paragraph()

# Дефект 2
doc.add_heading("Дефект №2. Экспресс-доставка на короткой дистанции выполняется за 0 дней", level=2)

t2 = doc.add_table(rows=6, cols=2)
t2.style = "Light Grid Accent 1"
cells = t2.rows[0].cells
cells[0].text = "Параметр"
cells[1].text = "Значение"

rows_data_2 = [
    ("Упавший тест", "test_express_on_short_distance_cannot_be_same_day"),
    ("Номер строки", "44 (файл delivery_service.py)"),
    ("Код строки", "if is_express: days_needed = days_needed // 2"),
    ("Описание дефекта",
     "После деления срока на 2 не применяется нижняя граница. Для дистанции "
     "1..499 км days_needed = max(1, 0) = 1, затем 1 // 2 = 0 дней, и дата "
     "доставки совпадает с датой отправки (2026-09-03). Даже экспресс не может "
     "доставить посылку мгновенно при дистанции больше 0."),
    ("Вариант исправления",
     "Оставить минимум 1 день после деления: "
     "days_needed = max(1, days_needed // 2)."),
]
for i, (k, v) in enumerate(rows_data_2, start=1):
    cells = t2.rows[i].cells
    cells[0].text = k
    cells[1].text = v

doc.add_paragraph()

# --- Результаты собственного проекта ----------------------------------
doc.add_heading("Исправление бага в собственном проекте", level=1)
doc.add_paragraph(
    "Тесты собственного проекта (my_project.py) изначально выявили 1 баг: "
    "функция mask_password() возвращала объект hashlib.sha256 вместо строки "
    "(отсутствовал вызов .hexdigest()). Это нарушало требование маскирования "
    "паролей: в лог попадал бы не хэш-отпечаток, а ссылка на объект хэша."
)
doc.add_paragraph("Исправление: mask_password() теперь возвращает sha256(password).hexdigest().")
doc.add_paragraph(
    "После исправления все 28 тестов собственного проекта проходят успешно."
)

doc.add_paragraph()

# --- Как запускать ----------------------------------------------------
doc.add_heading("Запуск тестов", level=1)
doc.add_paragraph("Из корня проекта (Lab2_FIO):")
doc.add_paragraph("python -m unittest discover -v   # полный прогон", style="Intense Quote")
doc.add_paragraph(
    "Результат итогового прогона: Ran 54 tests, FAILED (failures=2). "
    "Оба падения — ожидаемые и связаны с дефектами учебного модуля доставки."
)

OUT = "Отчет_ЛР2_АгиенкоДА.docx"
doc.save(OUT)
print(f"Сохранено: {OUT}")